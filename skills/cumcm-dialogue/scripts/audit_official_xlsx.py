#!/usr/bin/env python3
"""Read-only audit of a filled official workbook; never edits either file.

This verifies structure and allowed edits, not time mapping, annual coverage,
model correctness, or numerical optimality. Formula caches need separate review.
"""
from __future__ import annotations

import argparse
from copy import copy
import json
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries

ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
PLACEHOLDERS = {"...", "…", "⋮", "⁝", "︙"}


def _xml(obj: Any) -> str | None:
    if obj is None:
        return None
    if hasattr(obj, "to_tree"):
        return tostring(obj.to_tree(), encoding="unicode")
    return json.dumps(dict(obj), sort_keys=True)


def _cell_style(cell: Any) -> tuple:
    return (
        copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment),
        copy(cell.protection), cell.number_format,
    )


def audit(template: Path | str, result: Path | str,
          allowed_ranges: dict[str, list[str]] | None = None) -> dict[str, Any]:
    template = Path(template)
    result = Path(result)
    if template.resolve() == result.resolve():
        raise ValueError("Template and result must be different files")
    if template.suffix.lower() != ".xlsx" or result.suffix.lower() != ".xlsx":
        raise ValueError("This helper supports .xlsx only")
    books = []
    errors: list[dict[str, Any]] = []
    reviews: list[dict[str, Any]] = []
    filled = 0
    formulas = 0

    def error(code: str, location: str, detail: str) -> None:
        errors.append({"code": code, "location": location, "detail": detail})

    try:
        base = load_workbook(template, data_only=False)
        books.append(base)
        out = load_workbook(result, data_only=False)
        books.append(out)
        cached = load_workbook(result, data_only=True)
        books.append(cached)
        ranges: dict[str, list[tuple[int, int, int, int]]] = {}
        if allowed_ranges is not None:
            if not isinstance(allowed_ranges, dict):
                raise ValueError("Allowed ranges must be a worksheet-to-ranges object")
            for sheet, addresses in allowed_ranges.items():
                if sheet not in base.sheetnames or not isinstance(addresses, list):
                    raise ValueError(f"Invalid worksheet or ranges: {sheet}")
                ranges[sheet] = []
                for address in addresses:
                    bounds = range_boundaries(address)
                    if None in bounds or any(n < 1 for n in bounds):
                        raise ValueError(f"Invalid cell range: {sheet}!{address}")
                    c1, r1, c2, r2 = bounds
                    if c2 < c1 or r2 < r1 or r2 > base[sheet].max_row or c2 > base[sheet].max_column:
                        raise ValueError(f"Range extends outside template: {sheet}!{address}")
                    ranges[sheet].append(bounds)

        if base.sheetnames != out.sheetnames:
            error("sheetnames", "workbook", "Worksheet names or order changed")

        for bws in base.worksheets:
            if bws.title not in out.sheetnames:
                continue
            ows = out[bws.title]
            if (bws.max_row, bws.max_column) != (ows.max_row, ows.max_column):
                error("dimensions", bws.title,
                      f"Template {bws.max_row}x{bws.max_column}; result {ows.max_row}x{ows.max_column}")
            if sorted(map(str, bws.merged_cells.ranges)) != sorted(map(str, ows.merged_cells.ranges)):
                error("merged_cells", bws.title, "Merged ranges changed")
            if bws.freeze_panes != ows.freeze_panes:
                error("freeze_panes", bws.title, "Freeze panes changed")
            if bws.sheet_state != ows.sheet_state:
                error("sheet_state", bws.title, "Worksheet visibility changed")
            for key in ("sheet_format", "sheet_properties", "page_setup", "page_margins",
                        "print_options", "auto_filter", "data_validations"):
                if _xml(getattr(bws, key)) != _xml(getattr(ows, key)):
                    error(key, bws.title, f"{key} changed")
            for key in ("print_area", "print_title_rows", "print_title_cols"):
                if str(getattr(bws, key)) != str(getattr(ows, key)):
                    error(key, bws.title, f"{key} changed")
            for key in ("row_dimensions", "column_dimensions"):
                before = getattr(bws, key)
                after = getattr(ows, key)
                if set(before) != set(after) or any(_xml(before[k]) != _xml(after[k]) for k in before if k in after):
                    error(key, bws.title, "Explicit row/column formatting changed")

            for row in bws.iter_rows():
                for bc in row:
                    oc = ows[bc.coordinate]
                    loc = f"{bws.title}!{bc.coordinate}"
                    if isinstance(bc, MergedCell) != isinstance(oc, MergedCell):
                        error("merged_cell_type", loc, "Merged cell footprint changed")
                        continue
                    if _cell_style(bc) != _cell_style(oc):
                        error("cell_style", loc, "Original cell style changed")
                    if isinstance(bc, MergedCell):
                        continue
                    value_changed = bc.value != oc.value
                    permitted = any(c1 <= bc.column <= c2 and r1 <= bc.row <= r2
                                    for c1, r1, c2, r2 in ranges.get(bws.title, []))
                    if value_changed:
                        if allowed_ranges is not None and not permitted:
                            error("outside_fill_range", loc, "Value changed outside declared results")
                        elif bc.value is not None and not permitted:
                            error("fixed_value", loc, "Original nonempty value changed")
                        if bc.value is None and oc.value is not None:
                            filled += 1
                    if isinstance(bc.value, str) and bc.value.strip() in PLACEHOLDERS:
                        reviews.append({"code": "placeholder", "location": loc,
                                        "detail": "Check that sample/ellipsis rows can represent required coverage"})
                    if oc.data_type == "e" or (isinstance(oc.value, str) and oc.value in ERROR_VALUES):
                        error("excel_error", loc, str(oc.value))
                    if oc.data_type == "f":
                        formulas += 1
                        cv = cached[bws.title][bc.coordinate]
                        if cv.value is None:
                            reviews.append({"code": "uncached_formula", "location": loc,
                                            "detail": "Recalculate and verify the saved formula value"})
                        elif cv.data_type == "e" or (isinstance(cv.value, str) and cv.value in ERROR_VALUES):
                            error("cached_excel_error", loc, str(cv.value))
        if filled == 0:
            reviews.append({"code": "no_new_values", "location": "workbook",
                            "detail": "No new values; check whether this workbook actually contains results"})
        return {
            "status": "failed" if errors else "needs_review" if reviews else "passed",
            "checks_scope": "Structure and cell-edit rules only; numerical mapping and coverage are not verified",
            "template": template.name, "result": result.name,
            "filled_blank_cells": filled, "formulas": formulas,
            "errors": errors, "review_items": reviews,
        }
    finally:
        for book in books:
            book.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    parser.add_argument("result", type=Path)
    parser.add_argument("--allowed-ranges", type=Path)
    args = parser.parse_args()
    try:
        ranges = json.loads(args.allowed_ranges.read_text(encoding="utf-8")) if args.allowed_ranges else None
        report = audit(args.template, args.result, ranges)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({"status": "failed", "error": str(exc)}, ensure_ascii=False, indent=2))
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["errors"] else 2 if report["review_items"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
