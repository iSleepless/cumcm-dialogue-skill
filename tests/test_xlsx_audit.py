"""Behavior tests on generated workbooks; no competition or user data."""
import importlib.util
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

SCRIPT = Path(__file__).resolve().parents[1] / "skills/cumcm-dialogue/scripts/audit_official_xlsx.py"
spec = importlib.util.spec_from_file_location("xlsx_audit", SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class WorkbookAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name) / "template.xlsx"
        self.out = Path(self.tmp.name) / "result.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "plan"
        ws.append(["Date", "Energy (kWh)", "Cost (yuan)"])
        ws.append(["2025-02-01", None, None])
        ws.append(["2025-02-02", None, None])
        ws.append(["note", "official fixed text", None])
        ws.merge_cells("B4:C4")
        ws.column_dimensions["B"].width = 20
        ws.row_dimensions[1].height = 25
        ws["B2"].number_format = "0.00"
        ws.freeze_panes = "B2"
        wb.save(self.base)
        wb.close()

    def fill(self, mutate=None):
        wb = load_workbook(self.base)
        ws = wb["plan"]
        ws["B2"] = 123.4
        ws["C2"] = 51.2
        ws["B3"] = 0
        ws["C3"] = 0
        if mutate:
            mutate(wb, ws)
        wb.save(self.out)
        wb.close()

    def test_fill_without_structural_changes(self):
        self.fill()
        report = module.audit(self.base, self.out, {"plan": ["B2:C3"]})
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["filled_blank_cells"], 4)

    def test_header_change_is_rejected(self):
        self.fill(lambda wb, ws: setattr(ws["B1"], "value", "Power (kW)"))
        self.assertIn("fixed_value", {e["code"] for e in module.audit(self.base, self.out)["errors"]})

    def test_added_row_and_sheet_are_rejected(self):
        def change(wb, ws):
            ws.append(["2025-02-03", 5, 1])
            wb.create_sheet("extra")
        self.fill(change)
        codes = {e["code"] for e in module.audit(self.base, self.out)["errors"]}
        self.assertTrue({"dimensions", "sheetnames"}.issubset(codes))

    def test_restricted_fill_range_is_enforced(self):
        self.fill()
        report = module.audit(self.base, self.out, {"plan": ["B2:C2"]})
        self.assertIn("outside_fill_range", {e["code"] for e in report["errors"]})

    def test_merged_header_and_style_are_checked(self):
        def change(wb, ws):
            ws.unmerge_cells("B4:C4")
            ws["B2"].font = Font(size=17)
        self.fill(change)
        codes = {e["code"] for e in module.audit(self.base, self.out)["errors"]}
        self.assertTrue({"merged_cells", "cell_style"}.issubset(codes))

    def test_missing_formula_cache_is_not_called_verified(self):
        self.fill(lambda wb, ws: setattr(ws["C2"], "value", "=B2*0.4"))
        report = module.audit(self.base, self.out)
        self.assertEqual(report["status"], "needs_review")
        self.assertIn("uncached_formula", {i["code"] for i in report["review_items"]})

    def test_ellipsis_requires_coverage_review(self):
        wb = load_workbook(self.base)
        wb["plan"]["A3"] = "⁝"
        wb.save(self.base)
        wb.close()
        self.fill()
        report = module.audit(self.base, self.out)
        self.assertEqual(report["status"], "needs_review")
        self.assertIn("placeholder", {i["code"] for i in report["review_items"]})

    def test_excel_error_is_rejected(self):
        self.fill(lambda wb, ws: setattr(ws["C2"], "value", "#REF!"))
        report = module.audit(self.base, self.out)
        self.assertIn("excel_error", {e["code"] for e in report["errors"]})

    def test_column_width_change_is_rejected(self):
        self.fill(lambda wb, ws: setattr(ws.column_dimensions["B"], "width", 35))
        report = module.audit(self.base, self.out)
        self.assertIn("column_dimensions", {e["code"] for e in report["errors"]})

    def test_no_edit_side_effects(self):
        self.fill()
        before = (self.base.read_bytes(), self.out.read_bytes())
        module.audit(self.base, self.out)
        self.assertEqual(before, (self.base.read_bytes(), self.out.read_bytes()))


if __name__ == "__main__":
    unittest.main()
